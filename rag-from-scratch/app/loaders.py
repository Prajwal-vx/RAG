import zipfile
from pathlib import Path
from langchain_core.documents import Document
from pypdf import PdfReader
from docx import Document as DocxDocument
from app.config import settings

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx"}

EXECUTABLE_SIGNATURES = (
    b"MZ",            # DOS / Windows PE
    b"\x7fELF",       # Linux ELF
    b"\xca\xfe\xba\xbe", # Java Class / Mach-O fat binary
    b"\xfe\xed\xfa\xce", # Mach-O 32-bit
    b"\xfe\xed\xfa\xcf", # Mach-O 64-bit
)

def validate_file_signature(path: Path, ext: str) -> None:
    with path.open("rb") as f:
        header = f.read(1024)

    if not header:
        raise ValueError("Uploaded file is empty.")

    # Check for executable payloads
    for sig in EXECUTABLE_SIGNATURES:
        if header.startswith(sig):
            raise ValueError("Executable binaries are not permitted.")

    if ext == ".pdf":
        if not header.startswith(b"%PDF-"):
            raise ValueError("Invalid PDF file signature.")

    elif ext == ".docx":
        if not header.startswith(b"PK\x03\x04"):
            raise ValueError("Invalid DOCX file signature.")
        # Check ZIP safety and zip bomb prevention
        try:
            with zipfile.ZipFile(path, "r") as zf:
                total_uncompressed = 0
                max_allowed = settings.max_upload_size_mb * 1024 * 1024
                for info in zf.infolist():
                    # Check for zip directory traversal
                    if ".." in info.filename or info.filename.startswith(("/", "\\")):
                        raise ValueError("DOCX archive contains unsafe file paths.")
                    total_uncompressed += info.file_size
                    if total_uncompressed > max_allowed:
                        raise ValueError("DOCX archive exceeds safe uncompressed size limit.")
        except zipfile.BadZipFile:
            raise ValueError("Corrupt or invalid DOCX archive.")

    elif ext in {".txt", ".md"}:
        if b"\x00" in header:
            raise ValueError("Binary files are not permitted for plain text / markdown.")

def load_file(path: Path, document_id: str, original_filename: str | None = None) -> list[Document]:
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext}")

    validate_file_signature(path, ext)
    display_name = original_filename or path.name

    if ext == ".pdf":
        try:
            reader = PdfReader(str(path))
            if reader.is_encrypted:
                try:
                    reader.decrypt("")
                except Exception:
                    raise ValueError("Password-protected PDFs are not supported.")
            docs = []
            for page_no, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                if text.strip():
                    docs.append(Document(
                        page_content=text.strip(),
                        metadata={
                            "document_id": document_id,
                            "filename": display_name,
                            "page": page_no,
                        },
                    ))
            return docs
        except Exception as exc:
            if isinstance(exc, ValueError):
                raise
            raise ValueError(f"Failed to parse PDF document: {str(exc)}")

    elif ext == ".docx":
        try:
            docx = DocxDocument(str(path))
            text = "\n".join(p.text.strip() for p in docx.paragraphs if p.text.strip())
        except Exception as exc:
            raise ValueError(f"Failed to parse DOCX document: {str(exc)}")
    elif ext in {".txt", ".md"}:
        try:
            text = path.read_text(encoding="utf-8", errors="ignore").strip()
        except Exception as exc:
            raise ValueError(f"Failed to read text document: {str(exc)}")
    else:
        text = path.read_text(encoding="utf-8", errors="ignore").strip()

    if not text:
        return []

    return [Document(
        page_content=text,
        metadata={"document_id": document_id, "filename": display_name, "page": 1},
    )]
